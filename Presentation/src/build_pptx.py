"""PowerPoint version of the thesis discussion, built with python-pptx from the thesis tables and figures.

Usage: python build_pptx.py            -> ../powerpoint/thesis_presentation.pptx
       python build_pptx.py --single N -> one-slide deck for visual checks
Run `python3 make_assets.py` first (chart images). Needs python-pptx, Pillow, pdflatex and macOS
Quick Look (qlmanage), which renders the thesis figures and the formulas to PNG.
"""
import argparse
import re
import subprocess
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

from thesis_data import IMPUTERS, LEVEL_LABELS, PATTERN_NAMES, PATTERNS, THESIS_FIGURES, rows, series

HERE = Path(__file__).resolve().parent
OUT_DIR = HERE.parent / "powerpoint"
ASSETS = OUT_DIR / "assets"

TITLE_FONT, BODY_FONT = "Georgia", "Arial"
INK, ACCENT, MUTED = RGBColor(0x14, 0x21, 0x3D), RGBColor(0xE4, 0x57, 0x2E), RGBColor(0x6B, 0x72, 0x80)
TEXT, TINT, RULE = RGBColor(0x1F, 0x29, 0x37), RGBColor(0xF2, 0xF4, 0xF7), RGBColor(0xD0, 0xD5, 0xDD)
WHITE, PALE = RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xC9, 0xD3, 0xE3)
EM_ORANGE = RGBColor(0xFF, 0x95, 0x00)
SLIDE_W, SLIDE_H = 13.333, 7.5
MAIN_SLIDES = 15
AUTHOR = "Michele Carlo Fanelli"
FOOTER = "Michele Carlo Fanelli  ·  Bocconi University"
MOTIF = [(0, 6), (0, 8), (0, 10), (1, 3), (1, 5), (2, 1), (2, 2), (2, 3), (3, 3), (3, 8), (4, 3), (4, 6),
         (4, 14), (5, 2), (5, 11), (6, 10), (6, 11), (7, 2), (7, 3), (7, 4), (7, 6), (7, 8), (7, 9),
         (7, 11), (7, 12), (7, 13), (8, 1), (8, 2), (8, 6), (8, 11), (9, 8), (9, 13)]


# ------------------------------------------------------------------ assets
def render_pdf_png(pdf, out, size):
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["qlmanage", "-t", "-s", str(size), "-o", tmp, str(pdf)],
                       check=True, capture_output=True)
        img = Image.open(Path(tmp) / (Path(pdf).name + ".png")).convert("RGBA")
    flat = Image.new("RGB", img.size, "white")
    flat.paste(img, mask=img.split()[3])
    return flat


def thesis_figure(name):
    out = ASSETS / f"{name}.png"
    if not out.exists():
        render_pdf_png(THESIS_FIGURES / f"{name}.pdf", out, 3200).save(out)
    return out


FORMULAS = {
    "em": r"\mathbb{E}\left[r_{t,m}\mid r_{t,o}\right] = \bm{\mu}_m + \bm{\Sigma}_{mo}\,\bm{\Sigma}_{oo}^{-1}\left(r_{t,o}-\bm{\mu}_o\right)",
    "tangency": r"\hat w = \frac{\hat{\bm\Sigma}^{-1}\hat{\bm\mu}}{\mathbf{1}^{\top}\hat{\bm\Sigma}^{-1}\hat{\bm\mu}}",
    "dk": r"D_k = \frac{1}{\sqrt{k}}\Big(\sum_{i=1}^{k}\sin^2\theta_i\Big)^{1/2}\in[0,1]",
    "coverr": r"\mathrm{CovErr} = \lVert\hat{\bm\Sigma}-\bm\Sigma\rVert_F\,/\,\lVert\bm\Sigma\rVert_F",
    "inverse": r"\hat{\bm\Sigma}^{-1} = \textstyle\sum_i \hat\lambda_i^{-1}\,\hat v_i\hat v_i^{\top}",
}


def formula(name, color=(0x14, 0x21, 0x3D)):
    """LaTeX formula rendered to a transparent PNG; returns (path, height in points at 20 pt type)."""
    out = ASSETS / f"formula_{name}.png"
    px_per_pt = 3200 / (16 / 2.54 * 72)
    if out.exists() and out.stat().st_mtime > Path(__file__).stat().st_mtime:
        return out, Image.open(out).size[1] / px_per_pt
    tex = (r"\documentclass{article}\usepackage[paperwidth=16cm,paperheight=4cm,margin=4mm]{geometry}"
           r"\usepackage{amsmath,amssymb,bm}\pagestyle{empty}\begin{document}\centering"
           r"\fontsize{20}{24}\selectfont$\displaystyle " + FORMULAS[name] + r"$\end{document}")
    with tempfile.TemporaryDirectory() as tmp:
        (Path(tmp) / "f.tex").write_text(tex)
        subprocess.run(["pdflatex", "-interaction=nonstopmode", "f.tex"], cwd=tmp, check=True,
                       capture_output=True)
        img = render_pdf_png(Path(tmp) / "f.pdf", None, 3200).convert("L")
    alpha = img.point(lambda v: 255 - v)
    box = alpha.getbbox()
    pad = 12
    alpha = alpha.crop((box[0] - pad, box[1] - pad, box[2] + pad, box[3] + pad))
    rgba = Image.new("RGBA", alpha.size, color + (0,))
    rgba.putalpha(alpha)
    rgba.save(out)
    return out, alpha.size[1] / px_per_pt


# ------------------------------------------------------------------ text helpers
TOKEN = re.compile(r"(\*\*.+?\*\*|!!.+?!!|~~.+?~~|_\{.+?\}|\^\{.+?\}|\*[^*\s][^*]*?\*)")


def runs_of(text):
    out = []
    for part in TOKEN.split(text):
        if not part:
            continue
        outer = {"**": {"bold": True}, "!!": {"bold": True, "color": ACCENT}, "~~": {"color": MUTED}}
        if part[:2] in outer:
            out += [(t, {**outer[part[:2]], **st}) for t, st in runs_of(part[2:-2])]
        elif part.startswith("_{"):
            out.append((part[2:-1], {"baseline": "-25000"}))
        elif part.startswith("^{"):
            out.append((part[2:-1], {"baseline": "30000"}))
        elif part.startswith("*") and len(part) > 2:
            out.append((part[1:-1], {"italic": True}))
        else:
            out.append((part, {}))
    return out


def set_bullet(paragraph, color=ACCENT, indent=0.22, char="▪"):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(indent)))
    pPr.set("indent", str(-Inches(indent)))
    for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buChar", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    clr = OxmlElement("a:buClr")
    srgb = OxmlElement("a:srgbClr")
    srgb.set("val", str(color))
    clr.append(srgb)
    sz = OxmlElement("a:buSzPct")
    sz.set("val", "100000")
    font = OxmlElement("a:buFont")
    font.set("typeface", "Arial")
    ch = OxmlElement("a:buChar")
    ch.set("char", char)
    for el in (clr, sz, font, ch):
        pPr.append(el)


def text(slide, x, y, w, h, paras, size=14, color=TEXT, font=BODY_FONT, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=6, line_spacing=None,
         char_spacing=None):
    """Text box with no padding. `paras` is a string or a list of strings / dicts with markup."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        spec = para if isinstance(para, dict) else {"text": para}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = spec.get("align", align)
        p.space_after = Pt(spec.get("space_after", space_after))
        if spec.get("space_before"):
            p.space_before = Pt(spec["space_before"])
        if spec.get("line_spacing", line_spacing):
            p.line_spacing = spec.get("line_spacing", line_spacing)
        if spec.get("bullet"):
            set_bullet(p, spec.get("bullet_color", ACCENT), spec.get("indent", 0.22))
        p._p.get_or_add_endParaRPr().set("sz", str(int(spec.get("size", size) * 100)))
        for chunk, style in runs_of(spec["text"]):
            r = p.add_run()
            r.text = chunk
            f = r.font
            f.name = spec.get("font", font)
            f.size = Pt(spec.get("size", size))
            f.bold = style.get("bold", spec.get("bold", bold))
            f.italic = style.get("italic", spec.get("italic", False))
            f.color.rgb = style.get("color", spec.get("color", color))
            rPr = r._r.get_or_add_rPr()
            if "baseline" in style:
                rPr.set("baseline", style["baseline"])
            cs = spec.get("char_spacing", char_spacing)
            if cs:
                rPr.set("spc", str(int(cs * 100)))
    return box


def shape(slide, kind, x, y, w, h, fill, radius=None, alpha=None, line=None):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
        if alpha is not None:
            srgb = shp.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            a = OxmlElement("a:alpha")
            a.set("val", str(int(alpha * 100000)))
            srgb.append(a)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line[0]
        shp.line.width = Pt(line[1])
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius / min(w, h)
    return shp


def card(slide, x, y, w, h, fill=TINT):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, radius=0.08)


def picture(slide, path, x, y, w=None, h=None):
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                    Inches(w) if w else None, Inches(h) if h else None)


def formula_pic(slide, name, x, y, pt_size, center_w=None, color=(0x14, 0x21, 0x3D)):
    """Place a formula so that its type matches `pt_size`; optionally centre it in a width."""
    path, h_pt = formula(name, color)
    h_in = h_pt * pt_size / 20 / 72
    img = Image.open(path)
    w_in = h_in * img.size[0] / img.size[1]
    if center_w is not None:
        x = x + (center_w - w_in) / 2
    picture(slide, path, x, y, h=h_in)
    return w_in, h_in


def motif(slide, x, y, size=0.085, gap=0.04, colors=None):
    colors = colors or [RGBColor(0xC4, 0xC9, 0xD2), RGBColor(0xC4, 0xC9, 0xD2), ACCENT]
    for k, c in enumerate(colors):
        shape(slide, MSO_SHAPE.RECTANGLE, x + k * (size + gap), y, size, size, c)


def chrome(slide, number, kicker, title, dark=False):
    fg = WHITE if dark else INK
    motif(slide, 0.6, 0.47)
    text(slide, 1.0, 0.405, 9.0, 0.3, kicker, size=11, bold=True, color=ACCENT, char_spacing=1.5)
    text(slide, 0.6, 0.74, 12.13, 0.7, title, size=26, bold=True, color=fg, font=TITLE_FONT)
    foot = PALE if dark else MUTED
    text(slide, 0.6, 7.02, 6.0, 0.25, FOOTER, size=10, color=foot)
    text(slide, 11.23, 7.02, 1.5, 0.25, f"{number} / {MAIN_SLIDES}", size=10, color=foot,
         align=PP_ALIGN.RIGHT)


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body


def background(slide, color):
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H, color)


# ------------------------------------------------------------------ charts
def chart(slide, name, x, y, w, h):
    """Chart image drawn from the thesis tables by make_assets.py (rendered identically in every viewer)."""
    path = ASSETS / f"chart_{name}.png"
    if not path.exists():
        raise SystemExit(f"{path.name} missing: run `python3 make_assets.py` first")
    return picture(slide, path, x, y, w, h)


# ------------------------------------------------------------------ table helpers
NO_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def _border(cell, side, width_pt, color):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = {"L": "a:lnL", "R": "a:lnR", "T": "a:lnT", "B": "a:lnB"}[side]
    for el in tcPr.findall(qn(tag)):
        tcPr.remove(el)
    ln = OxmlElement(tag)
    ln.set("w", str(int(Pt(width_pt))))
    fill = OxmlElement("a:solidFill")
    clr = OxmlElement("a:srgbClr")
    clr.set("val", str(color))
    fill.append(clr)
    ln.append(fill)
    order = ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]
    idx = 0
    for i, child in enumerate(list(tcPr)):
        if child.tag in [qn(t) for t in order[:order.index(tag)]]:
            idx = i + 1
    tcPr.insert(idx, ln)


def table(slide, x, y, w, data, col_widths, row_h, size=10, header_rows=1, bold_mask=None,
          align=None, colors=None, rules=()):
    """Booktabs-style table: data is a list of rows of strings; rules lists row indices with a rule above."""
    nrows, ncols = len(data), len(data[0])
    gf = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h * nrows))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is None:
        style = OxmlElement("a:tableStyleId")
        tblPr.append(style)
    style.text = NO_STYLE
    for j, cw in enumerate(col_widths):
        tbl.columns[j].width = Inches(cw)
    for i in range(nrows):
        tbl.rows[i].height = Inches(row_h)
        for j in range(ncols):
            cell = tbl.cell(i, j)
            cell.fill.background()
            cell.margin_left = cell.margin_right = Inches(0.04)
            cell.margin_top = cell.margin_bottom = Inches(0)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = False
            p = tf.paragraphs[0]
            p.alignment = (align[j] if align else PP_ALIGN.LEFT)
            p._p.get_or_add_endParaRPr().set("sz", str(int(size * 100)))
            for chunk, st in runs_of(data[i][j]):
                r = p.add_run()
                r.text = chunk
                r.font.name = BODY_FONT
                r.font.size = Pt(size)
                is_bold = st.get("bold", False) or (bold_mask[i][j] if bold_mask else False) or i < header_rows
                r.font.bold = is_bold
                r.font.italic = st.get("italic", False)
                r.font.color.rgb = st.get("color", (colors[i][j] if colors else TEXT))
                if "baseline" in st:
                    r._r.get_or_add_rPr().set("baseline", st["baseline"])
    for j in range(ncols):
        _border(tbl.cell(0, j), "T", 1.25, INK)
        _border(tbl.cell(header_rows - 1, j), "B", 0.75, INK)
        _border(tbl.cell(nrows - 1, j), "B", 1.25, INK)
        for r_i in rules:
            _border(tbl.cell(r_i - 1, j), "B", 0.5, RULE)
    return tbl


# ------------------------------------------------------------------ slides
def slide_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, INK)
    cell, pitch, gx, gy = 0.22, 0.285, 8.25, 1.45
    for r in range(10):
        for c in range(16):
            if (r, c) in MOTIF:
                shape(s, MSO_SHAPE.RECTANGLE, gx + c * pitch, gy + r * pitch, cell, cell, ACCENT)
            else:
                shape(s, MSO_SHAPE.RECTANGLE, gx + c * pitch, gy + r * pitch, cell, cell, WHITE, alpha=0.1)
    text(s, 0.8, 1.45, 7.0, 0.3, "MASTER'S THESIS  ·  BOCCONI UNIVERSITY", size=12, bold=True,
         color=ACCENT, char_spacing=2)
    text(s, 0.8, 1.95, 7.2, 1.3, ["Imputation of Missing Data", "in Financial Time Series"], size=34,
         bold=True, color=WHITE, font=TITLE_FONT, space_after=0, line_spacing=1.05)
    text(s, 0.8, 3.5, 7.2, 0.45, "From reconstruction error to downstream decisions", size=19,
         color=PALE)
    for k, (label, value) in enumerate([("CANDIDATE", AUTHOR), ("SUPERVISOR", "Francesco Corielli"),
                                        ("ACADEMIC YEAR", "2025/2026")]):
        text(s, 0.8 + k * 3.55, 5.55, 3.4, 0.28, label, size=10, color=PALE, char_spacing=1.5)
        text(s, 0.8 + k * 3.55, 5.85, 3.4, 0.4, value, size=17, bold=True, color=WHITE)
    notes(s, "Introduce the problem in one sentence: financial panels always have holes, and the way "
             "they are filled changes what any model built on them concludes. The talk follows the "
             "thesis: motivation, design, results, then the two downstream evaluations.")


def slide_motivation(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, 2, "INTRODUCTION  ·  MOTIVATION", "Financial data is never fully observed")
    cols = [
        ("Where data goes missing", [
            "Order-book feeds drop quotes in bursts of cancellations",
            "Trading halts when prices move too violently",
            "OTC bond tenors go unquoted for days",
            "Macro data arrive at mixed frequencies, with lags"]),
        ("The usual fixes distort risk", [
            "Carrying the last price forward: stale prices bias betas ~~(Scholes and Williams, 1977)~~",
            "Dropping incomplete dates: unbiased only under MCAR, and much of the sample is lost",
            "Filling with a constant: variances and correlations are damped, even under MCAR"]),
        ("What is left open", [
            "Deep imputers assume MAR, usually implicitly through their masked training objective",
            "Evaluation is almost always pointwise, RMSE or MAE on masked cells, which says little "
            "about the covariance structure"]),
    ]
    for k, (head, items) in enumerate(cols):
        x = 0.6 + k * 4.14
        card(s, x, 1.7, 3.85, 3.65)
        text(s, x + 0.28, 1.95, 3.3, 0.4, head, size=17, bold=True, color=INK)
        text(s, x + 0.28, 2.5, 3.32, 2.75,
             [{"text": t, "bullet": True, "space_after": 9} for t in items], size=14.5)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 5.62, 12.13, 1.05, INK, radius=0.08)
    text(s, 0.95, 5.62, 11.5, 1.05,
         "**Aim of the thesis.**  Evaluate each imputation on its reconstruction accuracy and on the "
         "downstream model built on it, against the same model fitted on the complete panel.", size=18, color=WHITE,
         anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "Gaps are everywhere in finance. The default fixes are not neutral: stale prices bias betas, "
             "deletion throws away data, a constant fill damps correlations. The literature mostly scores "
             "imputations cell by cell; this thesis scores them by what a user builds on the result.")


RQS = [
    ("Mechanism or amount?", "What drives the damage, the share of missing data or the mechanism that "
                             "produced it, at comparable shares?"),
    ("Do the deep imputers earn their cost?", "Recurrent and diffusion models against classical likelihood "
                                              "and low-rank estimators, with training and sampling time counted."),
    ("Is pointwise accuracy a good proxy?", "Does a method that fills the holes accurately also preserve "
                                            "covariances, tails, marginals and temporal dependence?"),
    ("Do the rankings survive downstream?", "Is the ordering given by reconstruction metrics the ordering "
                                            "given by decisions taken on the imputed panel?"),
]


def _four_cards(s, entries, y0=1.7, h=1.72, body_size=14):
    for k, (head, body) in enumerate(entries):
        x = 0.6 + (k % 2) * 6.215
        y = y0 + (k // 2) * (h + 0.22)
        card(s, x, y, 5.915, h)
        text(s, x + 0.3, y + 0.2, 0.6, 0.8, str(k + 1), size=40, bold=True, color=ACCENT, font=TITLE_FONT)
        text(s, x + 1.0, y + 0.24, 4.7, 0.4, head, size=17, bold=True, color=INK)
        text(s, x + 1.0, y + 0.66, 4.65, h - 0.75, body, size=body_size)


def slide_questions(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, 3, "INTRODUCTION  ·  RESEARCH QUESTIONS", "Four research questions")
    _four_cards(s, RQS)
    text(s, 0.6, 5.62, 4.0, 0.25, "CONTRIBUTIONS", size=10, bold=True, color=MUTED, char_spacing=1.5)
    contribs = [
        ("Controlled testbed", "Fifty S&P 500 stocks, three mechanisms at five severities, seven imputers "
                               "in one framework."),
        ("Evaluation suite", "Past reconstruction error: covariance, Value-at-Risk, Wasserstein, "
                             "volatility clustering and leverage."),
        ("Downstream evaluations", "Tangency portfolio and factor structure, ranking imputers against "
                                   "each other and an oracle."),
    ]
    for k, (head, body) in enumerate(contribs):
        x = 0.6 + k * 4.14
        text(s, x, 5.92, 3.85, 0.95, [{"text": head, "bold": True, "color": ACCENT, "size": 14,
                                       "space_after": 3}, body], size=12.5)
    notes(s, "These four questions structure the results. Keep this slide short and come back to it at "
             "the end, where each question gets its answer.")


def slide_design(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, 4, "METHODOLOGY  ·  DATA AND DESIGN", "A controlled experiment on a complete equity panel")
    steps = [
        ("Complete panel", "S&P 500 end-of-day prices, 1/1/2010 to 30/5/2026. Stocks with any missing price "
                           "removed; 50 sampled across industries.", "4125 × 50 panel"),
        ("Inject missingness", "Three mechanisms, two MCAR and one MNAR, five levels each, matched on the "
                               "expected share of missing cells.", "15 corrupted panels"),
        ("Impute", "Seven methods, from deletion to conditional diffusion, run on every corrupted panel.",
         "105 runs (7 × 15)"),
        ("Evaluate", "Seven metrics against the true panel; two downstream models against an oracle "
                     "fitted on it.", "7 metrics + 2 models"),
    ]
    w, gap, y, h = 2.72, 0.417, 1.72, 2.95
    for k, (head, body, key) in enumerate(steps):
        x = 0.6 + k * (w + gap)
        card(s, x, y, w, h)
        shape(s, MSO_SHAPE.OVAL, x + 0.25, y + 0.25, 0.42, 0.42, ACCENT)
        text(s, x + 0.25, y + 0.25, 0.42, 0.42, str(k + 1), size=15, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.8, y + 0.29, w - 0.95, 0.4, head, size=15.5, bold=True, color=INK)
        text(s, x + 0.25, y + 0.88, w - 0.45, 1.45, body, size=12.5)
        text(s, x + 0.25, y + h - 0.55, w - 0.4, 0.35, key, size=15, bold=True, color=INK)
        if k < 3:
            shape(s, MSO_SHAPE.RIGHT_ARROW, x + w + 0.08, y + h / 2 - 0.13, gap - 0.16, 0.26,
                          RGBColor(0x9A, 0xA3, 0xB2))
    notes_text = [
        ("Returns, not prices", "Prices are positive, level dependent and non-stationary, so the panel is in daily log returns. "
                                "A deleted return is a stylised missing price: a carried-forward price is a "
                                "zero return, but the later catch-up return is not modelled."),
        ("Offline reconstruction", "Every method may use the full sample, including dates after the gap. "
                                   "The results describe reconstruction after the fact, not real-time filling."),
    ]
    for k, (head, body) in enumerate(notes_text):
        x = 0.6 + k * 6.215
        motif(s, x, 5.13, colors=[ACCENT])
        text(s, x + 0.22, 5.0, 5.6, 0.35, head, size=16, bold=True, color=INK)
        text(s, x + 0.22, 5.42, 5.65, 1.3, body, size=13.5)
    notes(s, "Start from a panel with no real missing data, delete entries under known rules, reconstruct, "
             "and compare with the truth. Returns rather than prices because prices are level dependent. "
             "Every method may use the full sample: this is reconstruction after the fact, not real-time "
             "filling.")


def slide_mechanisms(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, 5, "METHODOLOGY  ·  MISSINGNESS INJECTION",
           "Three mechanisms, matched on the share of missing cells")
    fig_w = 9.1
    fx = (SLIDE_W - fig_w) / 2
    picture(s, thesis_figure("missingness_examples"), fx, 1.42, w=fig_w)
    descs = [
        ("Sparse", "MCAR", "Each cell is deleted independently with probability *r*."),
        ("Contiguous gaps", "MCAR", "*n* seed cells, each extended forward into a run with mean length "
                                    "ℓ trading days."),
        ("Stressed", "MNAR", "Dates at or above the *p*-quantile of the squared panel-average return lose each "
                             "entry with probability *h*, other dates with *r*. The score is hidden from "
                             "the imputer."),
    ]
    col_w = fig_w / 3
    for k, (name, tag, body) in enumerate(descs):
        x = fx + k * col_w + 0.12
        text(s, x, 4.4, col_w - 0.3, 0.3, [{"text": f"**{name}**  ~~{tag}~~", "size": 14}], color=INK)
        text(s, x, 4.7, col_w - 0.25, 0.8, body, size=11.5)
    data = [["Expected share of missing cells", "", "0.1%", "1%", "5%", "10%", "30%"],
            ["Sparse", "*r*"] + LEVEL_LABELS["sparse"],
            ["Gaps", "(*n*, ℓ)"] + LEVEL_LABELS["gaps"],
            ["Stressed", "(*r*, *h*, *p*)"] + LEVEL_LABELS["stressed"]]
    colors = [[TEXT] * 7] + [[INK, MUTED] + [TEXT] * 5 for _ in range(3)]
    table(s, 0.6, 5.62, 12.13, data, [2.55, 1.1, 1.45, 1.75, 1.75, 1.75, 1.78], 0.29, size=11,
          align=[PP_ALIGN.LEFT, PP_ALIGN.LEFT] + [PP_ALIGN.CENTER] * 5, colors=colors)
    notes(s, "Sparse removes isolated cells, gaps remove runs within one asset, stressed removes entries "
             "preferentially on the most volatile dates, so deletions line up across assets. Levels are "
             "matched on the expected share, so any difference between patterns is due to the mechanism. "
             "The stressed score is never shown to the imputer, which makes the pattern non-ignorable.")


def slide_methods(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, 6, "METHODOLOGY  ·  MODELS", "Seven methods, from deletion to conditional diffusion")
    groups = [
        ("BASELINES", [("Deletion", "Drop every date with a missing value (complete cases)",
                        "Little and Rubin (2019)"),
                       ("Zero", "Fill with 0, as if the price were carried forward", "")]),
        ("CLASSICAL", [("KNN", "Inverse-distance average over the *k* nearest dates",
                        "Troyanskaya et al. (2001)"),
                       ("Low-rank", "Iterative truncated SVD: a *q*-factor reconstruction",
                        "Mazumder et al. (2010)"),
                       ("EM", "Gaussian maximum likelihood, conditional-mean fill", "Dempster et al. (1977)")]),
        ("DEEP LEARNING", [("BRITS", "Bidirectional RNN blending history and cross-section",
                            "Cao et al. (2018)"),
                           ("CSDI", "Conditional score-based diffusion; median of 32 draws",
                            "Tashiro et al. (2021)")]),
    ]
    for k, (head, methods) in enumerate(groups):
        x = 0.6 + k * 4.14
        card(s, x, 1.68, 3.85, 3.62)
        text(s, x + 0.28, 1.9, 3.3, 0.3, head, size=11, bold=True, color=ACCENT, char_spacing=1.5)
        paras = []
        for name, desc, ref in methods:
            paras.append({"text": name, "size": 16.5, "bold": True, "color": INK, "space_after": 1})
            paras.append({"text": desc, "size": 13, "space_after": 1})
            paras.append({"text": ref or " ", "size": 10.5, "color": MUTED, "space_after": 11})
        text(s, x + 0.28, 2.3, 3.32, 2.95, paras)
    card(s, 0.6, 5.5, 12.13, 1.3)
    formula_pic(s, "em", 0.6, 5.66, 24, center_w=12.13)
    text(s, 0.6, 6.38, 12.13, 0.3, "EM fill: the conditional mean given the assets observed on the same "
                                    "date, which minimises the expected squared error", size=12.5,
         color=MUTED, align=PP_ALIGN.CENTER)
    notes(s, "Two baselines that the literature expects to fail, three classical estimators and two deep "
             "models. EM is the one to watch: its fill is the conditional expectation given the other "
             "assets on the same date, the fill that minimises the expected squared error under the Gaussian model.")


def slide_evaluation(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, 7, "METHODOLOGY  ·  EVALUATION", "Evaluating the panel, not only the filled cells")
    text(s, 0.6, 1.62, 6.0, 0.35, "Seven metrics, four families", size=17, bold=True, color=INK)
    fams = [
        ("Pointwise", [("MAE, RMSE", "imputed cells only")]),
        ("Marginal", [("VaRErr", "5% Value-at-Risk, per asset"),
                      ("WassErr", "Wasserstein distance of marginals")]),
        ("Second moment", [("CovErr", "relative Frobenius error of \u03a3")]),
        ("Temporal", [("AbsACFErr", "volatility clustering, 20 lags"),
                      ("LevErr", "leverage, corr(*r*_{t}, |*r*_{t+k}|)")]),
    ]
    y = 2.12
    for fam, items in fams:
        h = 0.34 + 0.33 * len(items)
        card(s, 0.6, y, 6.15, h)
        text(s, 0.82, y + 0.17, 1.55, 0.3, fam, size=12.5, bold=True, color=ACCENT)
        text(s, 2.4, y + 0.16, 4.25, h - 0.2,
             [{"text": f"**{m}**   {d}", "space_after": 5} for m, d in items], size=12.5)
        y += h + 0.12
    text(s, 0.6, y + 0.05, 6.15, 0.5, "The first five are invariant to permuting the dates; the temporal "
                                      "ones are not.", size=12, color=MUTED)
    text(s, 7.0, 1.62, 5.7, 0.35, "Two downstream evaluations", size=17, bold=True, color=INK)
    card(s, 7.0, 2.12, 5.73, 1.72)
    text(s, 7.28, 2.3, 5.2, 0.35, "Tangency portfolio", size=15, bold=True, color=INK)
    w_in, _ = formula_pic(s, "tangency", 7.28, 2.72, 17)
    text(s, 7.28 + w_in + 0.2, 2.8, 12.5 - (7.28 + w_in + 0.2), 1.35, "fitted on each panel, evaluated on the true returns, compared with "
                                   "the oracle fitted on the complete panel (in sample)", size=12.5)
    card(s, 7.0, 4.02, 5.73, 2.65)
    text(s, 7.28, 4.2, 5.2, 0.35, "Factor structure", size=15, bold=True, color=INK)
    text(s, 7.28, 4.62, 5.2, 0.55, "Variance shares of the principal components and the subspace distance",
         size=12.5)
    formula_pic(s, "dk", 7.28, 5.25, 16)
    text(s, 7.28, 6.05, 5.2, 0.4, "between the true and estimated leading eigenspaces", size=12.5)
    notes(s, "Seven metrics in four families. Pointwise metrics look only at the filled cells, the others "
             "compare statistics of the whole panel, and the last two are the only ones that notice the "
             "order of the dates. The two downstream evaluations ask what a user would actually do with "
             "the panel.")


def slide_floor(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, 8, "RESULTS  ·  RECONSTRUCTION METRICS", "Pointwise errors hit a floor at once, and EM sits on it")
    chart(s, "mae_floor", 0.45, 1.6, 7.3, 4.95)
    text(s, 0.6, 6.6, 7.0, 0.3, "Sparse pattern, MAE on the imputed cells, by expected share of missing "
                                "cells (Table 3.1).", size=10.5, color=MUTED)
    text(s, 8.15, 1.75, 4.6, 0.6, "0.0071 to 0.0080", size=30, bold=True, color=ACCENT, font=TITLE_FONT)
    text(s, 8.15, 2.38, 4.6, 0.3, "EM's MAE across all five sparse levels", size=12.5, color=MUTED)
    text(s, 8.15, 2.95, 4.58, 3.9, [
        {"text": "Daily returns are close to serially uncorrelated: what can be known about a missing "
                 "entry is carried by the other assets on the same date.", "bullet": True},
        {"text": "EM's conditional mean extracts exactly that.", "bullet": True},
        {"text": "The five conditional methods are separated in the third decimal; the extra capacity of "
                 "BRITS and CSDI has little to act on.", "bullet": True},
        {"text": "The ranking is stable across patterns: EM is rarely beaten, and never by a wide margin.",
         "bullet": True},
    ], size=14.5, space_after=10)
    notes(s, "Daily returns carry almost no serial information, so the best one can do for a missing entry "
             "is to read it off the other assets on the same date. EM does exactly that, so its pointwise "
             "error is flat and it is very hard to beat. The deep models have nothing extra to learn here.")


def slide_structural(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, 9, "RESULTS  ·  POINTWISE VERSUS STRUCTURAL", "Accurate cells do not make an accurate panel")
    chart(s, "coverr_growth", 0.45, 1.55, 6.25, 3.75)
    chart(s, "zero_ratio", 6.88, 1.55, 5.85, 3.75)
    cards = [
        "MAE averages over the filled cells, CovErr accumulates them over the panel: EM's CovErr grows "
        "almost 40-fold while its MAE stays flat.",
        "At sparse 0.3, Zero's MAE is !!1.5 times!! the best method's and its CovErr !!8.9 times!! "
        "(ratios of the table values).",
        "The best fill cell by cell is a conditional mean, less variable than the return it replaces: "
        "Low-rank loses to EM on MAE but beats it on CovErr at two stressed levels.",
    ]
    for k, body in enumerate(cards):
        x = 0.6 + k * 4.14
        card(s, x, 5.45, 3.85, 1.38)
        text(s, x + 0.22, 5.55, 3.42, 1.22, body, size=12.5, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "Pointwise errors average over the filled cells, so they stay flat; the covariance error "
             "accumulates them, so it grows with the amount missing. Zero imputation shows the gap at its "
             "plainest: half again worse on MAE, nine times worse on covariance. Even among good methods the "
             "two families disagree: the best fill cell by cell damps the second moment most.")


def slide_mechanism(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, 10, "RESULTS  ·  MISSINGNESS PATTERNS", "The mechanism matters more than the amount")
    chart(s, "mechanism", 0.45, 1.6, 6.9, 4.95)
    text(s, 0.6, 6.6, 6.7, 0.3, "EM, covariance error by pattern at matched shares (Tables 3.1 to 3.3).",
         size=10.5, color=MUTED)
    # matched-share comparison
    card(s, 7.7, 1.68, 5.03, 1.85)
    text(s, 7.95, 1.82, 4.6, 0.3, "EM COVERR AT MATCHED SHARES", size=10.5, bold=True, color=MUTED,
         char_spacing=1.2)
    heads = ["", "Sparse", "Gaps", "Stressed"]
    vals = [("about 10%", ["0.0237", "0.0238", "0.0554"]), ("about 30%", ["0.0565", "0.0589", "0.6634"])]
    xs = [7.95, 9.35, 10.48, 11.6]
    for j, hd in enumerate(heads):
        text(s, xs[j], 2.17, 1.1, 0.3, hd, size=12, bold=True, color=INK)
    for i, (lab, v) in enumerate(vals):
        y = 2.55 + i * 0.42
        text(s, xs[0], y, 1.4, 0.35, lab, size=13, color=MUTED)
        for j, val in enumerate(v):
            hot = (i == 1 and j == 2)
            text(s, xs[j + 1], y - (0.05 if hot else 0), 1.15, 0.4, val, size=18 if hot else 14,
                 bold=hot, color=ACCENT if hot else TEXT, font=TITLE_FONT if hot else BODY_FONT)
    text(s, 7.7, 3.78, 5.03, 3.1, [
        {"text": "The two MCAR patterns are nearly interchangeable: with 49 other assets observed on each "
                 "date, breaking the time axis costs a cross-sectional imputer little.", "bullet": True},
        {"text": "Stressed deletions are aligned by date. They remove the conditioning set, and what they "
                 "remove are the largest returns.", "bullet": True},
        {"text": "BRITS falls behind where deletions are most extreme: MAE !!0.0301!! against EM's 0.0133 "
                 "at (0, 0.1, 0.99).", "bullet": True},
        {"text": "At (0.2, 1, 0.9) the ranking compresses: CovErr between 0.648 and 0.663 for every "
                 "surviving model.", "bullet": True},
    ], size=13, space_after=7)
    notes(s, "Read the tables across at matched shares. Sparse and gaps are almost the same, because on "
             "every date the other assets are still there. The stressed pattern removes exactly those "
             "dates, and the returns it removes are the largest ones, so the error sits in the target and "
             "no model capacity fixes it.")


def slide_portfolio(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, 11, "RESULTS  ·  DOWNSTREAM EVALUATION",
           "Downstream I: the portfolio amplifies the distortion")
    fig_w = 12.13
    picture(s, thesis_figure("portfolio_realizations"), 0.6, 1.5, w=fig_w)
    tiles = [
        ("Oracle", "Complete panel, annualised: mean 0.361, variance 0.049, Sharpe !!1.636!!."),
        ("Heavy MCAR", "Variance rises from 0.049 to as much as 0.29 and *L*^{1} weight distances reach "
                       "14: the inverted covariance over-leverages."),
        ("Heavy MNAR", "The imputers never see the volatility; their portfolios end at little more than half of the "
                       "oracle's cumulative return."),
        ("Who fails", "EM and Low-rank are never the worst. BRITS: Sharpe 1.30 against EM's 1.51 at "
                      "(0.05, 0.5, 0.9)."),
    ]
    w, gap = 2.87, 0.217
    for k, (head, body) in enumerate(tiles):
        x = 0.6 + k * (w + gap)
        card(s, x, 4.85, w, 1.65)
        text(s, x + 0.22, 4.98, w - 0.4, 0.32, head, size=15, bold=True, color=INK)
        text(s, x + 0.22, 5.35, w - 0.4, 1.1, body, size=12.5)
    text(s, 0.6, 6.6, 12.13, 0.35, "In-sample diagnostic of reconstruction fidelity, not investment "
                                   "performance. Deletion: Sharpe near zero at sparse 0.1; its covariance is "
                                   "singular at the heavy levels.", size=10.5, color=MUTED)
    notes(s, "The tangency portfolio inverts the covariance, so small distortions turn into large positions. "
             "Under random missingness the imputed portfolios take on far more risk than the oracle; under "
             "stressed missingness they never see the volatility and fall short. This is an in-sample "
             "diagnostic of fidelity, not a performance claim.")


def slide_factors(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, 12, "RESULTS  ·  DOWNSTREAM EVALUATION",
           "Downstream II: a robust market factor, a fragile subspace")
    fig_w = 7.75
    pic = picture(s, thesis_figure("factor_structure"), 0.55, 1.5, w=fig_w)
    fig_h = pic.height / 914400
    text(s, 0.6, 1.5 + fig_h + 0.06, 7.6, 0.5, "Top: variance share of the first 15 components (log "
                                               "scale). Bottom: ordered principal angles between the 15-factor "
                                               "subspaces, not per-factor errors. Heaviest level of each pattern.",
         size=10.5, color=MUTED)
    facts = [
        ("D_{1} < 0.026", "Market direction at sparse 0.3: under 1.5 degrees"),
        ("6 to 18×", "Larger D_{5}: the five-factor subspace, 0.14 to 0.39"),
        ("41.74% → 48.5%", "First-factor share at sparse 0.3: every conditional imputer raises it, to "
                            "between 45.6% and 48.5%"),
        ("10 of 15", "Levels where Low-rank has the largest D_{5}, although it imposes a factor structure"),
        ("≈ 25%", "First-factor share at (0.2, 1, 0.9), for all four methods that can be evaluated"),
    ]
    y = 1.55
    for big, small in facts:
        text(s, 8.65, y, 4.1, 0.42, big, size=21, bold=True, color=ACCENT, font=TITLE_FONT)
        text(s, 8.65, y + 0.42, 4.1, 0.5, small, size=12, color=TEXT)
        y += 1.04
    card(s, 0.6, 5.55, 7.7, 1.32)
    formula_pic(s, "inverse", 0.85, 5.66, 17)
    text(s, 0.85, 6.18, 7.25, 0.62,
         "The portfolio loads most on the trailing directions, the ones the imputations recover worst. Its "
         "error maximisation is the factor distortion seen through an inverse.", size=12)
    notes(s, "The first principal component is almost indestructible. The five-factor subspace is not, and "
             "the misalignment sits in one direction of it. Point imputations make the panel look more like "
             "its own factor model. Low-rank is worst on exactly the evaluation that might be expected to "
             "favour it. Last line: the portfolio inverse loads on the directions recovered worst, which "
             "ties the two evaluations together.")


def slide_cost(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, 13, "RESULTS  ·  COMPUTATIONAL COST", "The deep models did not earn their cost")
    tiles = [
        ("1 s", EM_ORANGE, "EM", "about one second to fill a 4125 × 50 panel"),
        ("1 h", INK, "BRITS", "just under an hour of training per panel (10,000 epochs)"),
        ("4 h", ACCENT, "CSDI", "just under four hours of training per panel, then about 24 minutes of "
                                "sampling"),
        ("70 h", INK, "All neural training", "roughly seventy hours for the 15 levels, four fifths of it "
                                             "on CSDI"),
    ]
    w, gap = 2.87, 0.217
    for k, (big, col, name, body) in enumerate(tiles):
        x = 0.6 + k * (w + gap)
        card(s, x, 1.75, w, 2.95)
        text(s, x, 1.95, w, 1.0, big, size=56, bold=True, color=col, font=TITLE_FONT, align=PP_ALIGN.CENTER)
        text(s, x + 0.2, 3.05, w - 0.4, 0.35, name, size=16, bold=True, color=INK, align=PP_ALIGN.CENTER)
        text(s, x + 0.25, 3.45, w - 0.5, 1.15, body, size=13, align=PP_ALIGN.CENTER)
    text(s, 0.6, 5.05, 12.13, 0.9, "CSDI matches EM's accuracy at about **four orders of magnitude** more "
                                   "computation, and BRITS falls behind at two stressed levels.", size=22,
         color=INK, font=TITLE_FONT)
    text(s, 0.6, 6.15, 12.13, 0.6, "Single Apple Mac mini (M4, 16 GB), CPU-only PyTorch. Small networks: "
                                   "about 120,000 parameters for BRITS and 47,000 for CSDI. "
                                   " Times are approximate, as reported in Section 3.3 of the thesis.",
         size=12, color=MUTED)
    notes(s, "All of this ran on a single consumer machine with small networks, so the costs come from the "
             "algorithms. EM needs about a second; CSDI needs hours of training and twenty-four minutes of "
             "sampling per panel, for a reconstruction that is not better.")


ANSWERS = [
    ("Mechanism or amount?", "**The mechanism, in this experiment.** Under MCAR the methods hold with 30% missing; "
                             "under stressed MNAR they break at far smaller shares, and at the heaviest "
                             "level no method separates."),
    ("Do the deep imputers earn their cost?", "**Not here.** EM is rarely beaten; CSDI matches it at four orders "
                                              "of magnitude more cost; BRITS falls behind at two stressed levels."),
    ("Is pointwise accuracy a good proxy?", "**Poor.** Zero's MAE is 1.5 times the best and its CovErr about 9 "
                                            "times. Conditional fills inflate the first factor, from 41.7% to as "
                                            "much as 48.5%, which no pointwise metric registers."),
    ("Do the rankings survive downstream?", "**Largely, and the exceptions are informative.** Low-rank: "
                                            "accurate cells, worst five-factor subspace at 10 of 15 levels. "
                                            "Deletion: inefficient in the tables, unstable portfolios downstream."),
]


def slide_answers(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, 14, "CONCLUSION  ·  SUMMARY", "Answers to the four questions")
    _four_cards(s, ANSWERS, y0=1.7, h=1.86, body_size=13.5)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 5.86, 12.13, 0.85, INK, radius=0.08)
    text(s, 0.6, 5.86, 12.13, 0.85, "Score an imputation against the use it is put to.", size=24,
         bold=True, color=WHITE, font=TITLE_FONT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "Go back to the four questions. The single message to leave with the committee is the last "
             "line: an imputation should be judged by the use it is put to, because reconstruction error is "
             "least informative exactly where the choice of method matters.")


def slide_limits(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, INK)
    chrome(s, 15, "CONCLUSION  ·  LIMITATIONS AND FUTURE WORK", "Limitations and future work", dark=True)
    cols = [
        ("LIMITATIONS", [
            "One panel: fifty liquid large caps, daily, one period, one run per configuration",
            "Injected missingness; the stressed pattern is a stylised MNAR that mixes the selection rule "
            "with date-aligned deletion",
            "Gaussian EM without the Student-*t* refinement; both downstream models read the covariance; "
            "the portfolio is in sample",
            "Offline setting: BRITS and CSDI use observations after the gap"]),
        ("FUTURE WORK", [
            "Causal imputation with only past data, to measure the cost of filling in real time",
            "Downstream tests beyond the second moment: VaR backtests, hedging, trading rules",
            "Multiple imputation: let the decision see the spread of the CSDI draws",
            "Larger cross sections, other frequencies and asset classes, real gaps"]),
    ]
    for k, (head, items) in enumerate(cols):
        x = 0.6 + k * 6.215
        text(s, x, 1.72, 5.6, 0.3, head, size=12, bold=True, color=ACCENT, char_spacing=1.5)
        text(s, x, 2.15, 5.65, 3.4, [{"text": t, "bullet": True, "space_after": 10} for t in items],
             size=15.5, color=WHITE)
    text(s, 0.6, 5.75, 8.0, 0.9, "Thank you", size=44, bold=True, color=WHITE, font=TITLE_FONT)
    notes(s, "Be upfront about the limits: one panel, one run, injected missingness, Gaussian EM, offline "
             "setting. The natural next step is the causal version of the problem, where only the past may "
             "be used.")


# ------------------------------------------------------------------ backup slides
def _backup_frame(prs, title):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s._element.set("show", "0")
    motif(s, 0.6, 0.43)
    text(s, 1.0, 0.36, 9.0, 0.3, "BACKUP", size=11, bold=True, color=ACCENT, char_spacing=1.5)
    text(s, 0.6, 0.64, 12.13, 0.45, title, size=17.5, bold=True, color=INK, font=TITLE_FONT)
    return s


def backup_metrics(prs, pattern, number, name):
    s = _backup_frame(prs, f"Evaluation scores under {name} missingness (Table {number}; best per level "
                           f"in bold)")
    head = ["Model", "MAE", "RMSE", "CovErr", "VaRErr", "WassErr", "AbsACFErr", "LevErr"]
    widths = [0.78] + [0.72] * 5 + [0.85, 0.72]
    align = [PP_ALIGN.LEFT] + [PP_ALIGN.RIGHT] * 7
    for side, levels in enumerate([(0, 1, 2), (3, 4)]):
        data, bold, colors, labels = [head], [[False] * 8], [[TEXT] * 8], []
        for r in rows("metrics", pattern):
            if r["level_index"] not in levels:
                continue
            if r["model"] == "Deletion":
                labels.append(len(data))
                data.append([f"Level {LEVEL_LABELS[pattern][r['level_index']]}"] + [""] * 7)
                bold.append([True] + [False] * 7)
                colors.append([ACCENT] + [TEXT] * 7)
            data.append([r["model"]] + ["--" if v is None else f"{v:.4f}" for v in r["values"]])
            bold.append([False] + r["bold"])
            colors.append([TEXT] * 8)
        tbl = table(s, 0.6 + side * 6.2, 1.15, sum(widths), data, widths, 0.235, size=8, bold_mask=bold,
                    align=align, colors=colors, rules=[i for i in labels if i > 1])
        for i in labels:
            tbl.cell(i, 0).merge(tbl.cell(i, 7))


def backup_matrix(prs, kind, blocks, title):
    s = _backup_frame(prs, title)
    columns = {"portfolio": ["Mean", "Variance", "Sharpe", "dw1"], "factors": ["PC1", "Top5", "D1", "D5"]}
    reference = "True (oracle)" if kind == "portfolio" else "True (reference)"
    head1 = [""] + sum([[PATTERN_NAMES[p], "", "", "", ""] for p in PATTERNS], [])
    head2 = ["Expected share"] + ["0.1%", "1%", "5%", "10%", "30%"] * 3
    data, bold, rules = [head1, head2], [[False] * 16, [False] * 16], []
    for column, label, fmt in blocks:
        j = columns[kind].index(column)
        ref = series(kind, "sparse", reference, column)[0]
        rules.append(len(data))
        data.append([f"**{label}**" + (f"  ~~(true {format(ref, fmt)})~~" if ref is not None else "")]
                    + [""] * 15)
        bold.append([False] * 16)
        for m in IMPUTERS:
            row, brow = [m], [False]
            for p in PATTERNS:
                for r in rows(kind, p):
                    if r["model"] == m:
                        v = r["values"][j]
                        row.append("--" if v is None else format(v, fmt))
                        brow.append(r["bold"][j])
            data.append(row)
            bold.append(brow)
    tbl = table(s, 0.6, 1.15, 12.13, data, [2.3] + [0.6553] * 15, 0.228, size=8, header_rows=2,
                bold_mask=bold, align=[PP_ALIGN.LEFT] + [PP_ALIGN.RIGHT] * 15, rules=rules)
    for k in range(3):
        tbl.cell(0, 1 + 5 * k).merge(tbl.cell(0, 5 + 5 * k))
        tbl.cell(0, 1 + 5 * k).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    for i in rules:
        tbl.cell(i, 0).merge(tbl.cell(i, 15))


# ------------------------------------------------------------------ build
SLIDES = [slide_title, slide_motivation, slide_questions, slide_design, slide_mechanisms, slide_methods,
          slide_evaluation, slide_floor, slide_structural, slide_mechanism, slide_portfolio, slide_factors,
          slide_cost, slide_answers, slide_limits,
          lambda prs: backup_metrics(prs, "sparse", "3.1", "sparse"),
          lambda prs: backup_metrics(prs, "gaps", "3.2", "contiguous-gap"),
          lambda prs: backup_metrics(prs, "stressed", "3.3", "stressed"),
          lambda prs: backup_matrix(prs, "portfolio",
                                    [("Sharpe", "Sharpe ratio", ".3f"),
                                     ("dw1", "L^{1} weight distance to the oracle", ".2f"),
                                     ("Mean", "Annualised mean", ".3f"),
                                     ("Variance", "Annualised variance", ".4f")],
                                    "Optimal portfolio at every level (Tables 3.4 to 3.6; best Sharpe in bold)"),
          lambda prs: backup_matrix(prs, "factors",
                                    [("PC1", "Variance share of the first component (%)", ".2f"),
                                     ("Top5", "Variance share of the first five (%)", ".2f"),
                                     ("D1", "Subspace distance D_{1}", ".4f"),
                                     ("D5", "Subspace distance D_{5}", ".4f")],
                                    "Factor structure at every level (Tables 3.7 to 3.9; best D_{5} in bold)")]


def register_notes_master(prs):
    """python-pptx relates the notes master but omits <p:notesMasterIdLst>, which Apple's importers require."""
    rid = next((r.rId for r in prs.part.rels.values() if r.reltype.endswith("/notesMaster")), None)
    root = prs.part._element
    if rid is None or root.find(qn("p:notesMasterIdLst")) is not None:
        return
    lst = OxmlElement("p:notesMasterIdLst")
    entry = OxmlElement("p:notesMasterId")
    entry.set(qn("r:id"), rid)
    lst.append(entry)
    root.find(qn("p:sldMasterIdLst")).addnext(lst)


def build(path, only=None):
    ASSETS.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    for k, make in enumerate(SLIDES, 1):
        if only is None or k == only:
            make(prs)
    if only is not None:
        for sld in prs.slides:
            sld._element.attrib.pop("show", None)
    register_notes_master(prs)
    prs.core_properties.title = "Imputation of Missing Data in Financial Time Series"
    prs.core_properties.author = AUTHOR
    prs.save(path)
    return path


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--single", type=int)
    ap.add_argument("--out")
    args = ap.parse_args()
    out = Path(args.out) if args.out else OUT_DIR / "thesis_presentation.pptx"
    print("wrote", build(out, args.single))
